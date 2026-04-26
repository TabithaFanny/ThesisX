import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


class DocxExporter:
    # Export Markdown to DOCX with Chinese academic paper formatting.

    HEADING_STYLES = {
        1: {
            "font": "\u9ed1\u4f53",
            "size": 18,
            "bold": True,
            "align": WD_ALIGN_PARAGRAPH.CENTER,
            "space_before": 12,
            "space_after": 6,
        },
        2: {
            "font": "\u9ed1\u4f53",
            "size": 15,
            "bold": True,
            "align": WD_ALIGN_PARAGRAPH.LEFT,
            "space_before": 12,
            "space_after": 6,
        },
        3: {
            "font": "\u9ed1\u4f53",
            "size": 14,
            "bold": True,
            "align": WD_ALIGN_PARAGRAPH.LEFT,
            "space_before": 6,
            "space_after": 6,
        },
        4: {
            "font": "\u9ed1\u4f53",
            "size": 12,
            "bold": True,
            "align": WD_ALIGN_PARAGRAPH.LEFT,
            "space_before": 6,
            "space_after": 3,
        },
        5: {
            "font": "\u9ed1\u4f53",
            "size": 12,
            "bold": True,
            "align": WD_ALIGN_PARAGRAPH.LEFT,
            "space_before": 3,
            "space_after": 3,
        },
        6: {
            "font": "\u9ed1\u4f53",
            "size": 12,
            "bold": False,
            "align": WD_ALIGN_PARAGRAPH.LEFT,
            "space_before": 3,
            "space_after": 3,
        },
    }
    BODY_FONT = "\u5b8b\u4f53"
    BODY_FONT_WESTERN = "Times New Roman"
    BODY_SIZE = 12
    LINE_SPACING = 1.5

    def export(self, markdown_text, output_path):
        doc = Document()
        self._setup_page(doc)
        self._setup_default_style(doc)
        blocks = self._parse_blocks(markdown_text)
        for block in blocks:
            self._render_block(doc, block)
        doc.save(output_path)

    def _setup_page(self, doc):
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    def _setup_default_style(self, doc):
        style = doc.styles["Normal"]
        font = style.font
        font.name = self.BODY_FONT_WESTERN
        font.size = Pt(self.BODY_SIZE)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), self.BODY_FONT)
        pf = style.paragraph_format
        pf.line_spacing = self.LINE_SPACING
        pf.first_line_indent = Pt(self.BODY_SIZE * 2)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

    def _parse_blocks(self, text):
        blocks = []
        lines = text.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i]

            if line.strip().startswith("```"):
                lang = line.strip()[3:].strip()
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1
                blocks.append(("code", "\n".join(code_lines), lang))
                continue

            m = re.match(r"^(#{1,6})\s+(.+)$", line)
            if m:
                level = len(m.group(1))
                title = re.sub(r"\s+#+\s*$", "", m.group(2).strip())
                blocks.append(("heading", title, level))
                i += 1
                continue

            if re.match(r"^\s*([-*_])\s*\1\s*\1[\s\1]*$", line):
                blocks.append(("hr",))
                i += 1
                continue

            if (
                "|" in line
                and i + 1 < len(lines)
                and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1])
            ):
                table_lines = []
                while i < len(lines) and "|" in lines[i]:
                    table_lines.append(lines[i])
                    i += 1
                blocks.append(("table", table_lines))
                continue

            if line.strip().startswith(">"):
                quote_lines = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    quote_lines.append(re.sub(r"^>\s?", "", lines[i]))
                    i += 1
                blocks.append(("quote", "\n".join(quote_lines)))
                continue

            m_ul = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
            if m_ul:
                list_items = []
                while i < len(lines):
                    m2 = re.match(r"^(\s*)([-*+])\s+(.*)$", lines[i])
                    if m2:
                        list_items.append(m2.group(3))
                        i += 1
                    else:
                        break
                blocks.append(("ul", list_items))
                continue

            m_ol = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
            if m_ol:
                list_items = []
                while i < len(lines):
                    m2 = re.match(r"^(\s*)(\d+)\.\s+(.*)$", lines[i])
                    if m2:
                        list_items.append(m2.group(3))
                        i += 1
                    else:
                        break
                blocks.append(("ol", list_items))
                continue

            if not line.strip():
                i += 1
                continue

            para_lines = []
            while i < len(lines) and lines[i].strip() and not self._is_block_start(lines[i]):
                para_lines.append(lines[i])
                i += 1
            if para_lines:
                blocks.append(("paragraph", " ".join(para_lines)))

        return blocks

    def _is_block_start(self, line):
        if re.match(r"^#{1,6}\s", line):
            return True
        if line.strip().startswith("```"):
            return True
        if re.match(r"^\s*([-*_])\s*\1\s*\1[\s\1]*$", line):
            return True
        if line.strip().startswith(">"):
            return True
        if re.match(r"^\s*[-*+]\s+", line):
            return True
        if re.match(r"^\s*\d+\.\s+", line):
            return True
        return False

    def _render_block(self, doc, block):
        btype = block[0]

        if btype == "heading":
            _, text, level = block
            p = doc.add_paragraph()
            style = self.HEADING_STYLES.get(level, self.HEADING_STYLES[6])
            p.alignment = style["align"]
            pf = p.paragraph_format
            pf.space_before = Pt(style["space_before"])
            pf.space_after = Pt(style["space_after"])
            pf.first_line_indent = Pt(0)
            pf.line_spacing = self.LINE_SPACING
            self._add_inline_text(
                p,
                text,
                default_font=style["font"],
                default_size=style["size"],
                default_bold=style["bold"],
            )

        elif btype == "paragraph":
            _, text = block
            p = doc.add_paragraph()
            self._add_inline_text(p, text)

        elif btype == "code":
            _, code, lang = block
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            pPr = p._element.get_or_add_pPr()
            shd = pPr.makeelement(qn("w:shd"), {qn("w:fill"): "F5F5F5", qn("w:val"): "clear"})
            pPr.append(shd)
            run = p.add_run(code)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        elif btype == "table":
            _, table_lines = block
            self._render_table(doc, table_lines)

        elif btype == "quote":
            _, text = block
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(1)
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = self.BODY_FONT_WESTERN
            run.font.size = Pt(self.BODY_SIZE)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            run.font.italic = True
            run.element.rPr.rFonts.set(qn("w:eastAsia"), self.BODY_FONT)

        elif btype == "ul":
            _, items = block
            for item in items:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.first_line_indent = Pt(0)
                self._add_inline_text(p, item)

        elif btype == "ol":
            _, items = block
            for item in items:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.first_line_indent = Pt(0)
                self._add_inline_text(p, item)

        elif btype == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            pPr = p._element.get_or_add_pPr()
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pBdr.makeelement(
                qn("w:bottom"),
                {
                    qn("w:val"): "single",
                    qn("w:sz"): "6",
                    qn("w:space"): "1",
                    qn("w:color"): "CCCCCC",
                },
            )
            pBdr.append(bottom)
            pPr.append(pBdr)

    def _render_table(self, doc, table_lines):
        rows_data = []
        for i, line in enumerate(table_lines):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if i == 1 and all(set(c.strip()) <= set("-: ") for c in cells):
                continue
            rows_data.append(cells)
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        for r in rows_data:
            while len(r) < num_cols:
                r.append("")

        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for r_idx, row_data in enumerate(rows_data):
            for c_idx, cell_text in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = ""
                p = cell.paragraphs[0]
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.space_before = Pt(2)
                pf.space_after = Pt(2)
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._add_inline_text(p, cell_text, default_size=10)
                if r_idx == 0:
                    for run in p.runs:
                        run.font.bold = True
                    tcPr = cell._element.get_or_add_tcPr()
                    shd = tcPr.makeelement(
                        qn("w:shd"), {qn("w:fill"): "F0F0F0", qn("w:val"): "clear"}
                    )
                    tcPr.append(shd)

    def _add_inline_text(
        self, paragraph, text, default_font=None, default_size=None, default_bold=None
    ):
        if default_font is None:
            default_font = self.BODY_FONT
        if default_size is None:
            default_size = self.BODY_SIZE
        if default_bold is None:
            default_bold = False

        tokens = self._tokenize_inline(text)
        for ttype, content, attrs in tokens:
            run = paragraph.add_run(content)
            run.font.name = self.BODY_FONT_WESTERN
            run.font.size = Pt(attrs.get("size", default_size))
            run.element.rPr.rFonts.set(qn("w:eastAsia"), attrs.get("east_font", default_font))
            if default_bold or attrs.get("bold"):
                run.font.bold = True
            if attrs.get("italic"):
                run.font.italic = True
            if attrs.get("underline"):
                run.font.underline = True
            if attrs.get("strike"):
                run.font.strike = True
            if attrs.get("color"):
                r, g, b = self._hex_to_rgb(attrs["color"])
                run.font.color.rgb = RGBColor(r, g, b)
            if attrs.get("highlight"):
                rPr = run.element.get_or_add_rPr()
                shd = rPr.makeelement(
                    qn("w:shd"),
                    {qn("w:fill"): attrs["highlight"].lstrip("#"), qn("w:val"): "clear"},
                )
                rPr.append(shd)
            if ttype == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(max(default_size - 1, 9))
                run.font.color.rgb = RGBColor(0xD6, 0x33, 0x84)
            if attrs.get("superscript"):
                run.font.superscript = True
            if attrs.get("subscript"):
                run.font.subscript = True

    def _tokenize_inline(self, text):
        tokens = []
        PATTERNS = [
            (r'<span\s+style="([^"]*)">(.*?)</span>', "styled"),
            (r"<sup>(.*?)</sup>", "superscript"),
            (r"<sub>(.*?)</sub>", "subscript"),
            (r"<u>(.*?)</u>", "underline"),
            (r"\*\*\*(.+?)\*\*\*", "bold_italic"),
            (r"\*\*(.+?)\*\*", "bold"),
            (r"__(.+?)__", "bold"),
            (r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", "italic"),
            (r"~~(.+?)~~", "strikethrough"),
            (r"`([^`]+)`", "code"),
            (r"\[([^\]]+)\]\([^\)]+\)", "link"),
            (r"!\[([^\]]*)\]\([^\)]+\)", "image"),
        ]

        combined = "|".join(f"({p})" for p, _ in PATTERNS)
        pattern_types = [t for _, t in PATTERNS]

        pos = 0
        for match in re.finditer(combined, text):
            if match.start() > pos:
                tokens.append(("plain", text[pos : match.start()], {}))

            full_match = match.group(0)
            matched_type = None
            inner = full_match
            # Explicitly capture styled groups before leaving the loop
            styled_style_str = ""
            styled_content_str = full_match

            for patt, ptype in PATTERNS:
                m2 = re.match(patt, full_match)
                if m2:
                    matched_type = ptype
                    inner = m2.group(2) if m2.lastindex and m2.lastindex >= 2 else m2.group(1)
                    if ptype == "styled":
                        styled_style_str = m2.group(1)
                        styled_content_str = (
                            m2.group(2) if m2.lastindex and m2.lastindex >= 2 else ""
                        )
                    break

            attrs = {}
            if matched_type == "bold":
                attrs["bold"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "italic":
                attrs["italic"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "bold_italic":
                attrs["bold"] = True
                attrs["italic"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "strikethrough":
                attrs["strike"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "code":
                tokens.append(("code", inner, attrs))
            elif matched_type == "underline":
                attrs["underline"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "superscript":
                attrs["superscript"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "subscript":
                attrs["subscript"] = True
                tokens.append(("text", inner, attrs))
            elif matched_type == "link":
                attrs["underline"] = True
                attrs["color"] = "#1a73e8"
                tokens.append(("text", inner, attrs))
            elif matched_type == "image":
                tokens.append(("text", f"[{inner}]", attrs))
            elif matched_type == "styled":
                attrs = self._parse_style(styled_style_str)
                tokens.append(("text", styled_content_str, attrs))
            else:
                tokens.append(("plain", full_match, {}))

            pos = match.end()

        if pos < len(text):
            tokens.append(("plain", text[pos:], {}))

        return tokens

    def _parse_style(self, style_str):
        attrs = {}
        m = re.search(r"(?<![a-z-])color:\s*([^;]+)", style_str)
        if m:
            attrs["color"] = m.group(1).strip()
        m = re.search(r"background-color:\s*([^;]+)", style_str)
        if m:
            attrs["highlight"] = m.group(1).strip()
        m = re.search(r"font-size:\s*(\d+)", style_str)
        if m:
            attrs["size"] = int(m.group(1))
        return attrs

    @staticmethod
    def _hex_to_rgb(color_str):
        color_str = color_str.strip().lstrip("#")
        if len(color_str) == 3:
            color_str = "".join(c * 2 for c in color_str)
        try:
            r = int(color_str[0:2], 16)
            g = int(color_str[2:4], 16)
            b = int(color_str[4:6], 16)
            return r, g, b
        except (ValueError, IndexError):
            return 0, 0, 0
