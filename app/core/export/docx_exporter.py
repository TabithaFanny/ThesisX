"""DocxExporter — Markdown to DOCX (Vision 3.9)."""

from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches


class DocxExporter:
    """Convert Markdown text to a .docx file preserving headings, bold, lists."""

    @staticmethod
    def export(markdown: str, output_path: Path, title: str = "论文") -> Path:
        doc = Document()

        # Title
        title_para = doc.add_heading(title, level=0)

        # Parse markdown line by line
        lines = markdown.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]

            # Heading
            heading_match = re.match(r"^(#{1,6})\s+(.+)", line)
            if heading_match:
                level = min(len(heading_match.group(1)), 6)
                doc.add_heading(heading_match.group(2).strip(), level=level)
                i += 1
                continue

            # Unordered list
            list_match = re.match(r"^[\s]*[-*+]\s+(.+)", line)
            if list_match:
                doc.add_paragraph(list_match.group(1), style="List Bullet")
                i += 1
                continue

            # Ordered list
            ordered_match = re.match(r"^[\s]*\d+\.\s+(.+)", line)
            if ordered_match:
                doc.add_paragraph(ordered_match.group(1), style="List Number")
                i += 1
                continue

            # Blockquote
            if line.startswith("> "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                run = p.add_run(line[2:])
                run.italic = True
                i += 1
                continue

            # Empty line → paragraph break
            if not line.strip():
                i += 1
                continue

            # Regular paragraph with inline bold
            p = doc.add_paragraph()
            _add_inline_runs(p, line)
            i += 1

        doc.save(str(output_path))
        return output_path


def _add_inline_runs(paragraph, text: str) -> None:
    """Handle **bold** and *italic* inline formatting."""
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)
